#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""記入済みの layout_input.xlsx（個別指定シート）を配列定義 JSON に変換する

配列マップ入力シートの内容は map_to_list.py で個別指定シートへ転記してから使います。

  python xlsx_to_json.py 新下駄.xlsx 新下駄.json
  python xlsx_to_json.py 新下駄.xlsx 新下駄.json でじこみ
"""

import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import json
import sys
from openpyxl import load_workbook

from typing_distance import (PHYSICAL_ROWS, THUMB_KEYS, SHIFT_KEYS,
                             keystrokes_kana, compute, GEOMETRIES)

VALID_KEYS = set(THUMB_KEYS) | set(SHIFT_KEYS)
for _row in PHYSICAL_ROWS.values():
    VALID_KEYS |= set(_row)


def clean(v):
    """セル値から全角スペース等を除去する（Excel入力でよく混入する）"""
    if v is None:
        return ""
    return str(v).replace("\u3000", "").replace("\xa0", "").strip()


def convert(xlsx_path, json_path, sample_text=None):
    wb = load_workbook(xlsx_path, data_only=True)

    if "個別指定" not in wb.sheetnames:
        raise SystemExit("個別指定シートが見つかりません。")

    top = wb["単打"] if "単打" in wb.sheetnames else wb["配列マップ入力"]

    name = clean(top["B5"].value)
    author = clean(top["B6"].value)
    if name.startswith("（ここに"):
        name = ""


    mapping = {}
    errors = []
    conflicts = []
    stats = {"単打": 0, "同時": 0, "連続": 0}

    # ---- 個別指定シート（唯一の正データ） ----
    if "個別指定" in wb.sheetnames:
        ws2 = wb["個別指定"]
        hrow = None
        for r in range(1, 15):
            if clean(ws2.cell(r, 1).value) == "かな" and clean(ws2.cell(r, 3).value) == "キー1":
                hrow = r
                break
        if hrow:
            for r in range(hrow + 1, ws2.max_row + 1):
                kana = ws2.cell(r, 1).value
                if not kana or str(kana).startswith("【例】"):
                    continue
                kana = clean(kana)
                ks = [clean(ws2.cell(r, c).value) for c in (3, 4, 5, 6)]
                ks = [k for k in ks if k]
                if not ks:
                    continue
                bad = [k for k in ks if k not in VALID_KEYS]
                if bad:
                    errors.append(f"個別指定 {r}行目 {kana}: 未知のキーID {', '.join(bad)}")
                    continue
                ktype = (ws2.cell(r, 7).value or "同時")
                ktype = str(ktype).strip()
                if len(ks) == 1:
                    spec = ks[0]
                    stats["単打"] += 1
                elif ktype == "連続":
                    spec = ks
                    stats["連続"] += 1
                else:
                    spec = "+".join(ks)
                    stats["同時"] += 1
                if kana in mapping and mapping[kana] != spec:
                    conflicts.append(f"{kana}: 重複行があります")
                mapping[kana] = spec

    if errors:
        print("[エラー]")
        for e in errors:
            print("  " + e)
        raise SystemExit("修正してから再実行してください")

    doc = {
        "name": name or "（無名）",
        "author": author,
        "input_mode": "kana_direct",
        "map": mapping,
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)

    print(f"{json_path} を出力しました")
    print(f"  配列名: {doc['name']}")
    print(f"  合計 {len(mapping)} 項目 "
          f"(単打 {stats['単打']} / 同時打鍵 {stats['同時']} / 連続打鍵 {stats['連続']})")
    if conflicts:
        print("  [注意] 同じかなが複数の位置に割り当てられています:")
        for c in conflicts[:10]:
            print("    " + c)

    if sample_text:
        st = keystrokes_kana(sample_text, doc)
        r = compute(st, GEOMETRIES["row_staggered"], layout_name=doc["name"])
        print(f"\n  お題文「{sample_text}」")
        print(f"  打鍵 {r.keys} / 動作 {r.actions} / 同時率 {r.keys_per_action:.2f}")
        print(f"  距離 左 {r.dist_L:.3f} / 右 {r.dist_R:.3f} / 合計 {r.total:.3f} u")
    return doc


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        raise SystemExit(1)
    convert(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
